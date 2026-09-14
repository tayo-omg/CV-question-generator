import os
import random
from io import BytesIO

from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from openai import OpenAI
from typing import List, Literal
from pydantic import BaseModel
from pypdf import PdfReader



load_dotenv()

api_key = os.getenv("OPENAI_API_KEY")

if not api_key:
    raise ValueError(
        "OPENAI_API_KEY was not found. Add it to your .env file."
    )

client = OpenAI(api_key=api_key)

app = FastAPI(
    title="CV Question",
    description=(
        "Extracts professional claims from an uploaded CV "
        "and generates 10 interview questions."
    ),
    version="1.0.0",
)


class CandidateQuestion(BaseModel):
    claim: str
    question: str
    expected_answer: List[str]
    difficulty: Literal["Easy", "Medium", "Hard"]
    red_flags: List[str]

class QuestionPool(BaseModel):
    claims: list[str]
    questions: list[CandidateQuestion]


class CVQuestionResponse(BaseModel):
    filename: str
    claims: list[str]
    questions: list[CandidateQuestion]


def extract_cv_text(
    file_content: bytes,
    filename: str,
) -> str:
    extension = os.path.splitext(filename.lower())[1]

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_content))

        pages = [
            page.extract_text() or ""
            for page in reader.pages
        ]

        return "\n".join(pages).strip()

    if extension == ".docx":
        document = Document(BytesIO(file_content))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_text = []

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        table_text.append(text)

        return "\n".join(
            paragraphs + table_text
        ).strip()

    raise ValueError(
        "Unsupported file type. Upload a PDF or DOCX file."
    )


def generate_question_pool(
    cv_text: str,
) -> QuestionPool:
    prompt = f"""
You are assisting a professional recruitment team.

Analyse the CV below and identify the candidate's main verifiable
professional claims.

Claims include:

- technical skills;
- projects completed;
- employment responsibilities;
- measurable achievements;
- experience using particular tools or technologies;
- leadership or problem-solving experience;
- educational qualifications and certifications.

Generate between 15 and 20 interview questions based only on claims
actually contained in the CV.

Requirements:

- Connect every question to a specific claim.
- Ask the candidate to explain how they performed the claimed work.
- Prefer practical, behavioural and scenario-based questions.
- Write one short and direct question for each claim.
- Each question must contain no more than 20 words.
- Ask only one main question at a time.
- Do not combine multiple questions into one.
- Focus on the most important part of the claim.
- Do not add a second sentence or follow-up question.
- Do not invent information.
- Do not ask about age, gender, religion, ethnicity, disability,
  marital status or other protected personal information.
- Do not generate general knowledge or trivia questions.
- Do not use contact information from the CV.

Example:

Claim: "Led a team that recovered debts exceeding N300 million
and reclaimed properties from debtors."

Good question:
"How did you lead your team to recover over N300 million in debt?"

Bad question:
"Describe the strategy used to recover these debts and reclaim
properties. How did you validate the claims, select legal or
negotiated remedies, allocate work across the team and measure
the amount recovered?"

- For every interview question, provide the following:

  1. difficulty:
     Classify the question as exactly one of:
     "Easy", "Medium", or "Hard".

     Use "Easy" for basic clarification of a skill, qualification or role, they should be able to explain quickly.
     Use "Medium" for questions requiring explanation of a process,
     responsibility, challenge or result.
     Use "Hard" for questions requiring detailed technical reasoning,
     leadership decisions, measurable evidence or complex problem-solving.

  2. expected_points:
     Provide between 2 and 4 short points that a strong answer should cover.
     Base every point only on the CV claim. Do not invent expected details
     that are not supported by the CV.


  3. red_flags:

     Expected terms they should be able to explain according to the claims in their CV
     Briefly state warning signs the interviewer should watch for in the
     candidate's answer. Include vague explanations, inability to describe
     personal contribution, inconsistent details, unsupported results,
     invented tools or failure to explain the process. Keep it concise and
     relevant to the specific CV claim.

     
CV:

{cv_text}
"""

    response = client.responses.parse(
        model=os.getenv("OPENAI_MODEL", "gpt-5.6"), input=prompt, text_format=QuestionPool, )

    result = response.output_parsed

    if result is None:
        raise ValueError(
            "The model did not return a valid response."
        )

    if len(result.questions) < 10:
        raise ValueError(
            "The model generated fewer than 10 questions."
        )

    return result


@app.get("/health")
def health_check():
    return {"status": "healthy", "service": "cv-question-api", }


@app.post(
    "/cv/questions",
    response_model=CVQuestionResponse,
)
async def create_cv_questions(
    cv: UploadFile = File(...),
):
    if not cv.filename:
        raise HTTPException(
            status_code=400, detail="The file must have a filename.",  )

    extension = os.path.splitext(
        cv.filename.lower()
    )[1]

    if extension not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=415, detail="Only PDF and DOCX files are supported.", )

    try:
        file_content = await cv.read()

        maximum_size = 5 * 1024 * 1024

        if not file_content:
            raise HTTPException(status_code=400, detail="The uploaded CV is empty.", )
          

        if len(file_content) > maximum_size:
            raise HTTPException( status_code=413,  detail="The CV must not exceed 5 MB.",  )

        cv_text = extract_cv_text(file_content=file_content, filename=cv.filename,  )

        if len(cv_text) < 100:
            raise HTTPException( status_code=422,
                detail=(
                    "Very little text was extracted. "
                    "The CV may be a scanned document."
                ),
            )

        question_pool = generate_question_pool(cv_text)

        selected_questions = random.SystemRandom().sample( question_pool.questions,  10,  )

        return CVQuestionResponse(
            filename=cv.filename, claims=question_pool.claims, questions=selected_questions,  )

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(status_code=500, detail=f"Question generation failed: {error}",
        ) from error

    finally:
        await cv.close()